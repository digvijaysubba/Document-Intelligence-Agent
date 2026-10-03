"""
Document Intelligence Agent
A Copilot-style document Q&A tool using Retrieval-Augmented Generation (RAG)
with Azure OpenAI, ChromaDB and SQLite
"""

import os
import sys
import sqlite3
import json
from datetime import datetime
from pathlib import Path
import PyPDF2
import chromadb
from docx import Document as DocxDocument
from openai import AzureOpenAI
from dotenv import load_dotenv

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent

missing = [name for name in ("AZURE_ENDPOINT", "AZURE_API_KEY") if not os.getenv(name)]
if missing:
    sys.exit(f"Missing settings: {', '.join(missing)}. Add them to {BASE_DIR / '.env'} (see README).")

AZURE_ENDPOINT = os.environ["AZURE_ENDPOINT"]
AZURE_API_KEY = os.environ["AZURE_API_KEY"]
DEPLOYMENT_NAME = os.getenv("AZURE_DEPLOYMENT_NAME", "gpt-4o")
EMBEDDING_DEPLOYMENT = os.getenv("AZURE_EMBEDDING_DEPLOYMENT", "text-embedding-3-small")
API_VERSION = os.getenv("AZURE_API_VERSION", "2024-12-01-preview")
HUMAN_REVIEW = os.getenv("HUMAN_REVIEW", "true").lower() != "false"

DB_PATH = BASE_DIR / "document_queries.db"
CHROMA_PATH = BASE_DIR / "chroma_db"
COLLECTION_NAME = "documents"
CHUNK_SIZE = 1000  # characters per chunk
CHUNK_OVERLAP = 200  # characters shared between neighbouring chunks
TOP_K = 4  # chunks retrieved per question
EMBEDDING_BATCH = 16  # chunks embedded per API call

client = AzureOpenAI(
    api_version=API_VERSION,
    azure_endpoint=AZURE_ENDPOINT,
    api_key=AZURE_API_KEY,
)


def init_database():
    """Initialize SQLite database for query history, upgrading older tables in place"""
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS queries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_name TEXT NOT NULL,
            query TEXT NOT NULL,
            response TEXT NOT NULL,
            timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """)

    existing = {row[1] for row in cursor.execute("PRAGMA table_info(queries)")}
    for column in ("status", "sources", "review_note"):
        if column not in existing:
            cursor.execute(f"ALTER TABLE queries ADD COLUMN {column} TEXT")

    conn.commit()
    conn.close()

def save_query(document_name, query, response, status, sources=None, review_note=None):
    """Save a query, its outcome and the chunks it was based on"""
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    cursor.execute("""
        INSERT INTO queries (document_name, query, response, timestamp, status, sources, review_note)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (document_name, query, response, datetime.now(), status,
          json.dumps(sources or []), review_note))

    conn.commit()
    conn.close()

def get_query_history(document_name=None, limit=10):
    """Retrieve query history from database"""
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    if document_name:
        cursor.execute("""
            SELECT document_name, query, response, timestamp, status FROM queries
            WHERE document_name = ?
            ORDER BY timestamp DESC LIMIT ?
        """, (document_name, limit))
    else:
        cursor.execute("""
            SELECT document_name, query, response, timestamp, status FROM queries
            ORDER BY timestamp DESC LIMIT ?
        """, (limit,))

    results = cursor.fetchall()
    conn.close()
    return results


def extract_text_from_pdf(file_path):
    """Extract text from PDF file"""
    text = ""
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += (page.extract_text() or "") + "\n"
    except Exception as e:
        print(f"Error extracting PDF: {e}")
    return text

def extract_text_from_docx(file_path):
    """Extract text from Word document, including tables"""
    text = ""
    try:
        doc = DocxDocument(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                text += " | ".join(cell.text for cell in row.cells) + "\n"
    except Exception as e:
        print(f"Error extracting DOCX: {e}")
    return text

def extract_text_from_txt(file_path):
    """Extract text from plain text file"""
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()
    except Exception as e:
        print(f"Error reading TXT: {e}")
    return ""

def load_document(file_path):
    """Load and extract text from any supported document type"""
    file_ext = Path(file_path).suffix.lower()

    if file_ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif file_ext == '.docx':
        return extract_text_from_docx(file_path)
    elif file_ext == '.txt':
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_ext}")


def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Split text into overlapping chunks, preferring to break at paragraph or line ends"""
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            # Break at the last blank line or newline in the second half of the window
            for separator in ("\n\n", "\n", ". "):
                cut = text.rfind(separator, start + chunk_size // 2, end)
                if cut != -1:
                    end = cut + len(separator)
                    break

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break
        start = end - overlap

    return chunks


def embed_texts(texts):
    """Turn a list of texts into embedding vectors with Azure OpenAI"""
    vectors = []
    for i in range(0, len(texts), EMBEDDING_BATCH):
        batch = texts[i:i + EMBEDDING_BATCH]
        response = client.embeddings.create(model=EMBEDDING_DEPLOYMENT, input=batch)
        vectors.extend(item.embedding for item in response.data)
    return vectors

def get_collection():
    """Open (or create) the persistent ChromaDB collection that stores document chunks"""
    chroma = chromadb.PersistentClient(path=str(CHROMA_PATH))
    return chroma.get_or_create_collection(
        name=COLLECTION_NAME,
        embedding_function=None,  # embeddings come from Azure OpenAI, not Chroma's default model
        metadata={"hnsw:space": "cosine"},
    )

def index_document(collection, document_name, text):
    """Chunk, embed and store a document, replacing any earlier copy of it"""
    chunks = chunk_text(text)
    if not chunks:
        return 0

    collection.delete(where={"document_name": document_name})
    collection.add(
        ids=[f"{document_name}::{i}" for i in range(len(chunks))],
        documents=chunks,
        embeddings=embed_texts(chunks),
        metadatas=[{"document_name": document_name, "chunk": i} for i in range(len(chunks))],
    )
    return len(chunks)

def retrieve_chunks(collection, question, document_name, top_k=TOP_K):
    """Semantic search: return the chunks of this document closest in meaning to the question"""
    available = len(collection.get(where={"document_name": document_name}, include=[])["ids"])
    if available == 0:
        return []

    results = collection.query(
        query_embeddings=embed_texts([question]),
        n_results=min(top_k, available),
        where={"document_name": document_name},
    )

    return [
        {"id": chunk_id, "text": text, "distance": round(distance, 3)}
        for chunk_id, text, distance in zip(
            results["ids"][0], results["documents"][0], results["distances"][0]
        )
    ]


def validate_prompt(query, document_summary):
    """
    Validate that the query is relevant to the document before answering.
    This is the input side of the 'responsible AI' layer.
    """
    validation_prompt = f"""
You are a query validator. Determine if the following query is relevant to a document.

Document Summary: {document_summary[:500]}

User Query: {query}

Respond with ONLY "VALID" or "INVALID" (no explanation).
A query is VALID if it's asking for information that could be found in the document.
A query is INVALID if it's completely unrelated or asking you to do something outside the document scope.
"""

    try:
        response = client.chat.completions.create(
            model=DEPLOYMENT_NAME,
            messages=[
                {"role": "system", "content": "You are a validation assistant. Respond with only VALID or INVALID."},
                {"role": "user", "content": validation_prompt}
            ],
            temperature=0.0,
            max_tokens=10
        )

        result = response.choices[0].message.content.strip().upper().strip('."')
        return result == "VALID"
    except Exception as e:
        print(f"Validation error: {e}")
        return False  # Fail closed: an unchecked question is not answered

def filter_output(answer, chunks, finish_reason):
    """
    Check a drafted answer before it is shown.
    This is the output side of the 'responsible AI' layer.
    Returns (passed, reason).
    """
    if finish_reason == "content_filter":
        return False, "Blocked by the Azure OpenAI content filter"

    if not answer or not answer.strip():
        return False, "Empty answer"

    context = "\n\n".join(f"[{i + 1}] {chunk['text']}" for i, chunk in enumerate(chunks))

    try:
        response = client.chat.completions.create(
            model=DEPLOYMENT_NAME,
            messages=[
                {"role": "system", "content": "You check answers for groundedness. Respond with only SUPPORTED or UNSUPPORTED."},
                {"role": "user", "content": f"""Context excerpts:
{context}

Answer:
{answer}

Is every factual claim in the answer supported by the context excerpts?
An answer that says the information is not in the document counts as SUPPORTED."""}
            ],
            temperature=0.0,
            max_tokens=10
        )

        verdict = response.choices[0].message.content.strip().upper().strip('."')
        if verdict != "SUPPORTED":
            return False, "Answer contains claims not supported by the retrieved excerpts"
        return True, "Grounded in retrieved excerpts"
    except Exception as e:
        print(f"Output filter error: {e}")
        return False, "Groundedness check failed"

def human_review(question, answer, chunks, filter_reason):
    """
    Human-in-the-loop step: a reviewer approves, edits or rejects the answer
    before it is surfaced. Returns (status, final_answer, note).
    """
    print("\n" + "-"*60)
    print("REVIEW REQUIRED")
    print("-"*60)
    print(f"Question: {question}")
    print(f"\nDraft answer:\n{answer}")
    print(f"\nAutomated check: {filter_reason}")
    print("\nSources:")
    for chunk in chunks:
        preview = " ".join(chunk["text"].split())[:150]
        print(f"  - {chunk['id']} (distance {chunk['distance']}): {preview}...")

    while True:
        decision = input("\n[a]pprove, [e]dit or [r]eject? ").strip().lower()

        if decision in ("a", "approve"):
            return "approved", answer, None
        if decision in ("e", "edit"):
            edited = input("Corrected answer: ").strip()
            if edited:
                return "edited", edited, "Answer edited by reviewer"
        if decision in ("r", "reject"):
            note = input("Reason (optional): ").strip()
            return "rejected", "This answer was rejected by the reviewer.", note or None


def summarize_document(text):
    """Generate a concise summary of the document"""
    content = text[:3000] if len(text) > 3000 else text

    try:
        response = client.chat.completions.create(
            model=DEPLOYMENT_NAME,
            messages=[
                {"role": "system", "content": "You are a document analyst. Provide concise, clear summaries."},
                {"role": "user", "content": f"Summarize this document in 2-3 sentences:\n\n{content}"}
            ],
            temperature=0.3,
            max_tokens=200
        )

        return response.choices[0].message.content
    except Exception as e:
        print(f"Summarization error: {e}")
        return "Unable to generate summary"


def extract_entities(text):
    """Extract key entities (names, dates, amounts, etc.) from document"""
    content = text[:4000] if len(text) > 4000 else text

    try:
        response = client.chat.completions.create(
            model=DEPLOYMENT_NAME,
            messages=[
                {"role": "system", "content": "You are an entity extraction expert. Extract and categorize key information."},
                {"role": "user", "content": f"""Extract key entities from this text and organize them as a JSON object with these keys:
- names (people/organizations)
- dates
- amounts (amounts/numbers)
- locations
- other_terms (other key terms)

Text:
{content}"""}
            ],
            temperature=0.0,
            max_tokens=600,
            response_format={"type": "json_object"}
        )

        response_text = response.choices[0].message.content
        try:
            return json.loads(response_text)
        except json.JSONDecodeError:
            return {"raw_extraction": response_text}
    except Exception as e:
        print(f"Entity extraction error: {e}")
        return {}


def answer_question(collection, question, document_name, summary):
    """Answer a question with RAG: validate, retrieve, generate, filter, review, log"""
    if not validate_prompt(question, summary):
        message = "Your question doesn't appear to be related to this document. Please ask something about the document content."
        save_query(document_name, question, message, "rejected_by_validation")
        return message

    try:
        chunks = retrieve_chunks(collection, question, document_name)
    except Exception as e:
        save_query(document_name, question, str(e), "error")
        return f"Error retrieving document context: {e}"

    if not chunks:
        message = "No indexed content was found for this document."
        save_query(document_name, question, message, "error")
        return message

    context = "\n\n".join(f"[{i + 1}] {chunk['text']}" for i, chunk in enumerate(chunks))
    sources = [chunk["id"] for chunk in chunks]

    try:
        response = client.chat.completions.create(
            model=DEPLOYMENT_NAME,
            messages=[
                {"role": "system", "content": "You are a helpful document assistant. Answer questions based ONLY on the provided excerpts. Cite excerpts as [1], [2], etc. If the answer is not in the excerpts, say so."},
                {"role": "user", "content": f"""Document excerpts:
{context}

Question: {question}

Provide a clear, concise answer based only on the excerpts above."""}
            ],
            temperature=0.2,
            max_tokens=500
        )

        answer = response.choices[0].message.content or ""
        finish_reason = response.choices[0].finish_reason
    except Exception as e:
        save_query(document_name, question, str(e), "error", sources)
        return f"Error generating answer: {e}"

    passed, filter_reason = filter_output(answer, chunks, finish_reason)

    if HUMAN_REVIEW:
        status, final_answer, note = human_review(question, answer, chunks, filter_reason)
    elif passed:
        status, final_answer, note = "auto_approved", answer, filter_reason
    else:
        status, final_answer, note = "blocked_by_filter", "The answer could not be verified against the document, so it was withheld.", filter_reason

    save_query(document_name, question, final_answer, status, sources, note)
    return final_answer


def run_interactive_session(file_path):
    """Run an interactive Q&A session with a document"""
    print("\n" + "="*60)
    print("DOCUMENT INTELLIGENCE AGENT")
    print("="*60)

    print(f"\nLoading document: {file_path}")
    try:
        document_text = load_document(file_path)
        if not document_text.strip():
            print("Error: Could not extract text from document")
            return
    except Exception as e:
        print(f"Error loading document: {e}")
        return

    document_name = Path(file_path).name

    print("\nIndexing document (chunking and embedding)...")
    collection = get_collection()
    try:
        chunk_count = index_document(collection, document_name, document_text)
    except Exception as e:
        print(f"Error indexing document: {e}")
        print(f"Check that an embedding deployment named '{EMBEDDING_DEPLOYMENT}' exists, or set AZURE_EMBEDDING_DEPLOYMENT.")
        return
    print(f"Indexed {chunk_count} chunks into ChromaDB")

    print("\nGenerating summary...")
    summary = summarize_document(document_text)
    print(f"\nSUMMARY:\n{summary}")

    print("\nExtracting key entities...")
    entities = extract_entities(document_text)
    print(f"\nKEY ENTITIES:\n{json.dumps(entities, indent=2)}")

    print("\n" + "-"*60)
    print("Ask questions about the document (type 'quit' to exit)")
    if HUMAN_REVIEW:
        print("Human review is ON: each answer must be approved before it is shown")
    print("-"*60)

    while True:
        question = input("\nYour question: ").strip()

        if question.lower() in ['quit', 'exit', 'q']:
            break

        if not question:
            continue

        print("\nThinking...")
        answer = answer_question(collection, question, document_name, summary)
        print(f"\nANSWER:\n{answer}")

    print("\n" + "-"*60)
    print("QUERY HISTORY:")
    print("-"*60)
    history = get_query_history(document_name, limit=5)
    for doc, q, a, ts, status in history:
        print(f"\n[{ts}] ({status or 'answered'})")
        print(f"Q: {q}")
        print(f"A: {a[:100]}...")


if __name__ == "__main__":
    init_database()

    if len(sys.argv) > 1:
        run_interactive_session(sys.argv[1])
        sys.exit()

    test_file = BASE_DIR / "sample_document.txt"

    if not test_file.exists():
        with open(test_file, 'w', encoding='utf-8') as f:
            f.write("""
SERVICE AGREEMENT

This Service Agreement ("Agreement") is entered into as of January 15, 2024,
between ABC Corporation ("Service Provider") and XYZ Industries ("Client").

1. SERVICES
The Service Provider agrees to provide software development and consulting services
as requested by the Client. Services include system design, implementation, testing,
and deployment.

2. PAYMENT TERMS
- Monthly retainer: $5,000
- Additional hourly work: $150 per hour
- Payment due within 30 days of invoice
- Late payment penalty: 2% per month

3. TERM AND TERMINATION
- Initial term: 12 months from the date above
- Either party may terminate with 30 days written notice
- Upon termination, all project deliverables must be transferred to Client

4. CONFIDENTIALITY
Both parties agree to maintain confidentiality of proprietary information
shared during the engagement. This includes source code, business strategies,
and client data.

5. LIABILITY
Service Provider's total liability shall not exceed the total fees paid in the
preceding 12 months. This excludes data breaches and gross negligence.
""")
        print(f"Created sample document: {test_file}")

    run_interactive_session(test_file)
